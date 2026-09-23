#!/usr/bin/env python3
"""
benchmark/real_world/scripts/build_prjna1088182_cohort.py

Parses Widney & Copley 2024 supplementary metadata (media-1.xlsx, media-2.xlsx)
and NCBI SRA runinfo to build the Phase D 25-sample cohort manifest and
publication truth dataset for post-edit genome audit validation.
"""

import csv
import openpyxl
from pathlib import Path

REPO_DIR = Path(__file__).resolve().parent.parent.parent.parent
METADATA_DIR = REPO_DIR / "benchmark" / "real_world" / "metadata"
MANIFEST_DIR = REPO_DIR / "benchmark" / "real_world" / "manifests"
TRUTH_DIR = REPO_DIR / "benchmark" / "real_world" / "truth"
RESULTS_DIR = REPO_DIR / "benchmark" / "real_world" / "results" / "prjna1088182"

MANIFEST_DIR.mkdir(parents=True, exist_ok=True)
TRUTH_DIR.mkdir(parents=True, exist_ok=True)
RESULTS_DIR.mkdir(parents=True, exist_ok=True)

# 1. Load SRA RunInfo
sra_by_biosample = {}
sra_by_sample_name = {}
sra_runinfo_path = MANIFEST_DIR / "prjna1088182_sra_runinfo.csv"

with open(sra_runinfo_path, "r", encoding="utf-8") as f:
    reader = csv.DictReader(f)
    for row in reader:
        avg_len = float(row["avgLength"]) if row["avgLength"] else 0
        bs_id = row["BioSample"].replace("SAMN", "")
        sample_name = row["SampleName"]
        if avg_len > 200:
            sra_by_biosample[bs_id] = row
            sra_by_sample_name[sample_name] = row

print(f"Loaded {len(sra_by_biosample)} PE SRA runs.")

# 2. Parse media-1.xlsx by exact column indices and aggregate by BioSample
wb = openpyxl.load_workbook(METADATA_DIR / "media-1.xlsx", data_only=True)
biosample_records = {}

def get_clean_str(val):
    if val is None:
        return ""
    s = str(val).strip()
    return "" if s.lower() in ["none", "n/a", "na", "-", "0"] else s

# Sheet 1: Cas9 edited strains
ws = wb["Cas9 edited strains"]
for r in list(ws.iter_rows(values_only=True))[2:]:
    if not any(c is not None for c in r):
        continue
    name, s_num, bs_link = r[0], r[1], str(r[4]) if r[4] else ""
    bs_id = bs_link.rstrip("/").split("/")[-1]
    if not bs_id:
        continue

    if bs_id not in biosample_records:
        proto = str(r[10]).strip() if r[10] else "."
        cut_site = str(r[8]).strip() if r[8] else "."
        gene = str(r[7]).strip() if r[7] else "."
        biosample_records[bs_id] = {
            "name": name,
            "sheet": "Cas9 edited strains",
            "group": "Cas9",
            "edited_gene": gene,
            "cut_site": cut_site,
            "protospacer": proto,
            "mutations": [],
        }

    # Mutation 1: pos 11, type 12, orig 13, new 14, gene 15
    p1 = get_clean_str(r[11])
    if p1:
        biosample_records[bs_id]["mutations"].append({
            "pos": p1.replace(",", ""),
            "type": str(r[12]).strip() if r[12] else "unspecified",
            "orig": str(r[13]).strip() if r[13] else ".",
            "new": str(r[14]).strip() if r[14] else ".",
            "gene": str(r[15]).strip() if r[15] else ".",
            "colony": str(s_num),
        })
    # Mutation 2: pos 17, type 18, orig 19, new 20, gene 21
    if len(r) > 17:
        p2 = get_clean_str(r[17])
        if p2:
            biosample_records[bs_id]["mutations"].append({
                "pos": p2.replace(",", ""),
                "type": str(r[18]).strip() if r[18] else "unspecified",
                "orig": str(r[19]).strip() if r[19] else ".",
                "new": str(r[20]).strip() if r[20] else ".",
                "gene": str(r[21]).strip() if r[21] else ".",
                "colony": str(s_num),
            })

# Sheet 2: lambda Red only strains
ws = wb["lambda Red only strains"]
for r in list(ws.iter_rows(values_only=True))[2:]:
    if not any(c is not None for c in r):
        continue
    name, s_num, bs_link = r[0], r[1], str(r[4]) if r[4] else ""
    bs_id = bs_link.rstrip("/").split("/")[-1]
    if not bs_id:
        continue

    if bs_id not in biosample_records:
        gene = str(r[7]).strip() if r[7] else "."
        biosample_records[bs_id] = {
            "name": name,
            "sheet": "lambda Red only strains",
            "group": "lambdaRed",
            "edited_gene": gene,
            "cut_site": ".",
            "protospacer": ".",
            "mutations": [],
        }

    # Mutation 1: pos 10, type 11, orig 12, new 13, gene 14
    p1 = get_clean_str(r[10])
    if p1:
        biosample_records[bs_id]["mutations"].append({
            "pos": p1.replace(",", ""),
            "type": str(r[11]).strip() if r[11] else "unspecified",
            "orig": str(r[12]).strip() if r[12] else ".",
            "new": str(r[13]).strip() if r[13] else ".",
            "gene": str(r[14]).strip() if r[14] else ".",
            "colony": str(s_num),
        })

# Sheet 3: I-SceI edited strains
ws = wb["I-SceI edited strains"]
for r in list(ws.iter_rows(values_only=True))[2:]:
    if not any(c is not None for c in r):
        continue
    name, s_num, bs_link = r[0], r[1], str(r[4]) if r[4] else ""
    bs_id = bs_link.rstrip("/").split("/")[-1]
    if not bs_id:
        continue

    if bs_id not in biosample_records:
        gene = str(r[7]).strip() if r[7] else "."
        biosample_records[bs_id] = {
            "name": name,
            "sheet": "I-SceI edited strains",
            "group": "I-SceI",
            "edited_gene": gene,
            "cut_site": ".",
            "protospacer": ".",
            "mutations": [],
        }

    p1 = get_clean_str(r[10])
    if p1:
        biosample_records[bs_id]["mutations"].append({
            "pos": p1.replace(",", ""),
            "type": str(r[11]).strip() if r[11] else "unspecified",
            "orig": str(r[12]).strip() if r[12] else ".",
            "new": str(r[13]).strip() if r[13] else ".",
            "gene": str(r[14]).strip() if r[14] else ".",
            "colony": str(s_num),
        })
    if len(r) > 16:
        p2 = get_clean_str(r[16])
        if p2:
            biosample_records[bs_id]["mutations"].append({
                "pos": p2.replace(",", ""),
                "type": str(r[17]).strip() if r[17] else "unspecified",
                "orig": str(r[18]).strip() if r[18] else ".",
                "new": str(r[19]).strip() if r[19] else ".",
                "gene": str(r[20]).strip() if r[20] else ".",
                "colony": str(s_num),
            })

# Sheet 4: control strains
ws = wb["control strains"]
for r in list(ws.iter_rows(values_only=True))[2:]:
    if not any(c is not None for c in r):
        continue
    name, s_num, bs_link = r[0], r[1], str(r[4]) if r[4] else ""
    bs_id = bs_link.rstrip("/").split("/")[-1]
    if not bs_id:
        continue

    if bs_id not in biosample_records:
        biosample_records[bs_id] = {
            "name": name,
            "sheet": "control strains",
            "group": "control",
            "edited_gene": ".",
            "cut_site": ".",
            "protospacer": ".",
            "mutations": [],
        }

    p1 = get_clean_str(r[7])
    if p1:
        biosample_records[bs_id]["mutations"].append({
            "pos": p1.replace(",", ""),
            "type": str(r[8]).strip() if r[8] else "unspecified",
            "orig": str(r[9]).strip() if r[9] else ".",
            "new": str(r[10]).strip() if r[10] else ".",
            "gene": str(r[11]).strip() if r[11] else ".",
            "colony": str(s_num),
        })
    if len(r) > 13:
        p2 = get_clean_str(r[13])
        if p2:
            biosample_records[bs_id]["mutations"].append({
                "pos": p2.replace(",", ""),
                "type": str(r[14]).strip() if r[14] else "unspecified",
                "orig": str(r[15]).strip() if r[15] else ".",
                "new": str(r[16]).strip() if r[16] else ".",
                "gene": str(r[17]).strip() if r[17] else ".",
                "colony": str(s_num),
            })

# 3. Target 25 samples
TARGET_SELECTIONS = {
    "Cas9": [
        "Cas9_1",             # single clone, mut: pos 301690 (paoB syn)
        "Cas9_2",             # single clone, clean
        "Cas9_103",           # single clone, clean
        "Cas9_3 _and2others", # mut: pos 1571257
        "Cas9_101_102",       # 2-colony pool, clean
        "Cas9_104_105",       # 2-colony pool, clean
        "Cas9_106_107",       # 2-colony pool, clean
        "Cas9_108_109",       # 2-colony pool, clean
        "Cas9_110_111",       # 2-colony pool, clean
        "Cas9_112_113",       # 2-colony pool, clean
    ],
    "control": [
        "control_1_2_3_4",     # clean (Reference starter)
        "control_5_6_7_8",     # mut: pos 1635887
        "control_9_10_11_12",   # clean
        "control_13_14_15_16", # clean
        "control_17_18_19_20", # mut: pos 3108123
    ],
    "lambdaRed": [
        "lambdaRed_1_2",   # clean
        "lambdaRed_3_4",   # mut: pos 3112908
        "lambdaRed_5_6",   # clean
        "lambdaRed_7_8",   # mut: pos 161068
        "lambdaRed_9_10",  # clean
    ],
    "I-SceI": [
        "I-SceI_1_2",        # mut: pos 2653047, 4115617
        "I-SceI_3_4_5_6",    # clean
        "I-SceI_7_8_9",      # mut: pos 2872949 (two identical colonies)
        "I-SceI_14_15_16",   # mut: pos 4068404
        "I-SceI_17_18_19",   # mut: pos 3073571
    ],
}

manifest_rows = []
truth_rows = []
control_starter_run = sra_by_sample_name["control_1_2_3_4"]["Run"]

for group, sample_names in TARGET_SELECTIONS.items():
    for s_name in sample_names:
        sra_row = sra_by_sample_name.get(s_name)
        if not sra_row:
            continue

        run_id = sra_row["Run"]
        bs_id = sra_row["BioSample"].replace("SAMN", "")
        meta = biosample_records.get(bs_id, {})

        editor = "cas9" if group == "Cas9" else "dsb"
        protospacer = meta.get("protospacer", ".")
        spacer = "."
        pam = "."
        if protospacer and protospacer != "." and len(str(protospacer)) >= 20:
            proto_str = str(protospacer).strip()
            if len(proto_str) >= 23:
                spacer = proto_str[:20]
                pam = proto_str[20:]
            else:
                spacer = proto_str

        sample_type = "single_clone" if s_name in ["Cas9_1", "Cas9_2", "Cas9_103"] else "pooled"
        starter_id = control_starter_run if run_id != control_starter_run else run_id

        r1_path = f"data/prjna1088182/{run_id}_1.fastq.gz"
        r2_path = f"data/prjna1088182/{run_id}_2.fastq.gz"
        ref_path = "data/references/MG1655.gbk"

        mut_list = meta.get("mutations", [])
        notes = f"{group} cohort ({s_name}): {meta.get('name', '')}"
        if mut_list:
            notes += f"; pub mutations={','.join(m['pos'] for m in mut_list)}"

        manifest_rows.append({
            "dataset_id": "prjna1088182",
            "sample_id": run_id,
            "sample_name": s_name,
            "group": group,
            "starter_id": starter_id,
            "r1": r1_path,
            "r2": r2_path,
            "reference": ref_path,
            "intended_tsv": ".",
            "editor": editor,
            "spacer": spacer,
            "pam": pam if pam != "." else "NGG",
            "truth_source": "PUBLICATION_REPORTED",
            "sample_type": sample_type,
            "notes": notes,
        })

        if mut_list:
            for m in mut_list:
                truth_rows.append({
                    "dataset": "prjna1088182",
                    "sample_id": run_id,
                    "sample_name": s_name,
                    "group": group,
                    "reported_mutation": m["pos"],
                    "pos": m["pos"],
                    "ref_allele": m["orig"],
                    "alt_allele": m["new"],
                    "mutation_type": m["type"],
                    "gene": m["gene"],
                    "edited_gene": meta.get("edited_gene", "."),
                    "cut_site": meta.get("cut_site", "."),
                })
        else:
            truth_rows.append({
                "dataset": "prjna1088182",
                "sample_id": run_id,
                "sample_name": s_name,
                "group": group,
                "reported_mutation": "NONE",
                "pos": ".",
                "ref_allele": ".",
                "alt_allele": ".",
                "mutation_type": "NONE",
                "gene": ".",
                "edited_gene": meta.get("edited_gene", "."),
                "cut_site": meta.get("cut_site", "."),
            })

# Write cohort manifest TSV
manifest_tsv = MANIFEST_DIR / "prjna1088182_cohort.tsv"
fieldnames = [
    "dataset_id", "sample_id", "sample_name", "group", "starter_id",
    "r1", "r2", "reference", "intended_tsv", "editor", "spacer", "pam",
    "truth_source", "sample_type", "notes"
]
with open(manifest_tsv, "w", encoding="utf-8", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter="\t")
    writer.writeheader()
    writer.writerows(manifest_rows)

print(f"Wrote cohort manifest: {manifest_tsv} ({len(manifest_rows)} samples)")

# Write truth TSV
truth_tsv = TRUTH_DIR / "prjna1088182_publication_truth.tsv"
truth_fields = [
    "dataset", "sample_id", "sample_name", "group", "reported_mutation",
    "pos", "ref_allele", "alt_allele", "mutation_type", "gene",
    "edited_gene", "cut_site"
]
with open(truth_tsv, "w", encoding="utf-8", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=truth_fields, delimiter="\t")
    writer.writeheader()
    writer.writerows(truth_rows)

print(f"Wrote publication truth: {truth_tsv} ({len(truth_rows)} entries)")
